from openpyxl import Workbook

# Create workbook
wb = Workbook()

# Select and rename active sheet
ws = wb.active
ws.title = "Tasks"

# Add headers
ws["A1"] = "Tasks"
ws["B1"] = "Status"

# List of Tasks
tasks = ["Learn Python", "Do Pushups", "Apply Freelance", "Build Project"]

# Add tasks in sheet
for i, task in enumerate(tasks, start=2):
    ws[f"A{i}"] = task
    ws[f"B{i}"] = "Pending"

# Save one task as completed
ws["B2"] = "Complete"

# Save workbook
wb.save("tasks.xlsx")

print("Excel file created successfully")
