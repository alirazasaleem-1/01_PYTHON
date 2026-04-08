import json
from openpyxl import Workbook, load_workbook
from datetime import datetime

# ================ Input ===================
def get_task_input():
    task = input("Enter Your Task: ")
    reminder = input("Enter Reminder Time (e.g: 6 PM ): ")

    date = datetime.now().strftime(r"%Y-%m-%d")

    data = {
    "Task": task,
    "Reminder": reminder,
    "Date": date
    }
    return data 

data = get_task_input() 
print("\n[LOG] ✅ Task Captured Successfully. ")

# ============== Save to JSON ===============
def save_to_json(data):
    try:
        with open(r"D:\01_PYTHON\DAY 58\tasks.json", 'r') as file:
            tasks_list = json.load(file)
    except:
            tasks_list = []

    tasks_list.append(data)

    with open(r"D:\01_PYTHON\DAY 58\tasks.json", 'w') as file:
        json.dump(tasks_list, file, indent=4)

save_to_json(data)
print("[LOG] ✅ Saved to JSON 📁")

# ============== Save to Excel ===============
def save_to_excel(data):
    try:
         wb = load_workbook(r"D:\01_PYTHON\DAY 58\tasks.xlsx")
         ws = wb.active
    except: 
         wb = Workbook()
         ws = wb.active
         ws.title = "Tasks"
         ws.append(["Task", "Reminder", "Date"]) # Headers

    ws.append([data["Task"], data["Reminder"], data["Date"]])
    wb.save(r"D:\01_PYTHON\DAY 58\tasks.xlsx")
save_to_excel(data)
print("[LOG] ✅ Saved to Excel 📊")

print("✅🎉 Task Added Successfully. ")