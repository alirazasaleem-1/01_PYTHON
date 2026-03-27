from openpyxl import Workbook, load_workbook
import os

file_name = r"D:\01_PYTHON\tasks.xlsx"

# Create file if not exists
if not os.path.exists(file_name):
    # Create workbook
    wb = Workbook()
    # Select Active Sheet
    ws = wb.active
    ws.title = ("Tasks")
    ws.append(["Task", "Status"])
    wb.save(file_name)

# Load file
wb = load_workbook(file_name)
ws = wb["Tasks"]

# Function to add task
def add_task(task):
    ws.append([task, "Pending"])
    wb.save(file_name)
    print("Task added Successfully. ✅")

# Function to remove task
def remove_task(task_name):
    for row in ws.iter_rows(min_row=2):
        if row[0].value == task_name:
            ws.delete_rows(row[0].row)
            wb.save(file_name)
            print("Task removed Succesfully. ✅")
            found = True
            break

    if not found:
        print("Task Not Found. ❌")
            

# Function to show pending tasks
def show_tasks():
    print("\nPending Tasks: ")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] == "Pending":
            print(f"- {row[0]}")

# MENU
while True:
    print("Welcome to Excel Task Manager")
    print("\n1. Add Task\n2. Remove Task\n3. Show Tasks\n4. Exit")
    choice = input("Enter Choice: ")


    if choice == "1":
        task = input("Enter the Task to Add: ")
        add_task(task)

    elif choice == "2":
        task = input("Enter the Task to Remove: ")
        remove_task(task)

    elif choice == "3":
        show_tasks()

    elif choice == "4":
        print("Good Bye 👋...!! ")
        break

    else:
        print("Invalid Choice. ")




