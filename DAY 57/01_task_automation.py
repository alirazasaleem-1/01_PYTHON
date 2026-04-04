import time
import json 
import os
from openpyxl import Workbook, load_workbook

print("Welcome to Ali's Task Manager")
task = input("Enter Your Task: ")
if task.strip() == "":
    print("❌ Task cannot be empty. ")
    exit()
current = time.strftime("%Y-%m-%d %H:%M:%S")

# -------------------------------------------------------------- Excel Part ------------------------------------------------------------

if os.path.exists(r"D:\01_PYTHON\DAY 57\Tasks.xlsx"):
    wb = load_workbook(r"D:\01_PYTHON\DAY 57\Tasks.xlsx")
    ws = wb.active
else:
    wb = Workbook()
    ws = wb.active
    ws.title ="Tasks"
    ws.append(["Task", "Timestamp"])

ws.append([task, current])
wb.save(r"D:\01_PYTHON\DAY 57\Tasks.xlsx")

# ---------------------------------------------------------------JSON Part -------------------------------------------------------------

tasks = []

if os.path.exists(r"D:\01_PYTHON\DAY 57\Tasks.json"):
    with open(r"D:\01_PYTHON\DAY 57\Tasks.json", "r") as file:
        try:
            tasks = json.load(file)
        except:
            tasks = []
        
tasks.append({
    "task": task,
    "time": current
})

with open(r"D:\01_PYTHON\DAY 57\Tasks.json", 'w') as file:
    json.dump(tasks, file, indent=4)

print("✅ Task Saved Successfully. ")